"""Tests for the Excel generation service."""

import sys
from pathlib import Path
from io import BytesIO

PROJECT_ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(PROJECT_ROOT))

import pytest
import openpyxl
from services.excel_service import generate_excel


@pytest.fixture
def sample_expenses():
    return [
        {"date": "2024-06-15", "description": "COTO", "amount": 12500.50, "category": "Supermercado", "source": "test.pdf"},
        {"date": "2024-06-16", "description": "UBER", "amount": 2300.00, "category": "Transporte", "source": "test.pdf"},
    ]


def test_generate_excel_returns_bytes(sample_expenses):
    """generate_excel must return raw bytes."""
    result = generate_excel(sample_expenses)
    assert isinstance(result, bytes)
    assert len(result) > 0


def test_generate_excel_creates_valid_workbook(sample_expenses):
    """The returned bytes must form a valid openpyxl workbook."""
    result = generate_excel(sample_expenses)
    wb = openpyxl.load_workbook(BytesIO(result))
    assert wb.active.title == "Gastos"


def test_generate_excel_headers(sample_expenses):
    """The first row must contain the expected Spanish headers."""
    result = generate_excel(sample_expenses)
    wb = openpyxl.load_workbook(BytesIO(result))
    ws = wb.active
    headers = [ws.cell(row=1, column=c).value for c in range(1, 6)]
    assert headers == ["Fecha", "Descripcion", "Monto", "Categoria", "Fuente"]


def test_generate_excel_data_rows(sample_expenses):
    """Data rows must match the input expenses exactly."""
    result = generate_excel(sample_expenses)
    wb = openpyxl.load_workbook(BytesIO(result))
    ws = wb.active

    assert ws.cell(row=2, column=1).value == "2024-06-15"
    assert ws.cell(row=2, column=2).value == "COTO"
    assert ws.cell(row=2, column=3).value == 12500.50
    assert ws.cell(row=2, column=4).value == "Supermercado"
    assert ws.cell(row=2, column=5).value == "test.pdf"


def test_generate_excel_empty_list():
    """An empty expense list should still produce a valid workbook with headers only."""
    result = generate_excel([])
    wb = openpyxl.load_workbook(BytesIO(result))
    ws = wb.active
    assert ws.max_row == 1  # only header row
    assert ws.cell(row=1, column=1).value == "Fecha"


def test_generate_excel_default_category():
    """Expenses without explicit category default to 'Otros'."""
    expenses = [{"date": "2024-06-20", "description": "X", "amount": 100}]
    result = generate_excel(expenses)
    wb = openpyxl.load_workbook(BytesIO(result))
    ws = wb.active
    assert ws.cell(row=2, column=4).value == "Otros"


def test_generate_excel_number_format(sample_expenses):
    """Amount cells must use a numeric format."""
    result = generate_excel(sample_expenses)
    wb = openpyxl.load_workbook(BytesIO(result))
    ws = wb.active
    monto_cell = ws.cell(row=2, column=3)
    assert monto_cell.number_format == '#,##0.00'
